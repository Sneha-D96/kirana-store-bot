import os
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

def generate_pdf_invoice(bill_id, bill_data, items):
    """Generates a branded, realistic, and monochrome invoice PDF."""
    os.makedirs("invoices", exist_ok=True)
    filename = f"invoices/invoice_{bill_id}.pdf"
    
    doc = SimpleDocTemplate(filename, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    elements = []
    styles = getSampleStyleSheet()
    
    # Updated to realistic dark/black styling
    title_style = ParagraphStyle('StoreTitle', parent=styles['Heading1'], fontSize=22, textColor=colors.HexColor("#111111"), alignment=1, fontName="Helvetica-Bold")
    subtitle_style = ParagraphStyle('StoreSub', parent=styles['Normal'], fontSize=9, textColor=colors.HexColor("#333333"), alignment=1)
    bold_style = ParagraphStyle('BoldText', parent=styles['Normal'], fontSize=10, fontName="Helvetica-Bold")
    
    elements.append(Paragraph("KIRANA STORE", title_style))
    elements.append(Paragraph("<b>GSTIN:</b> 33AAAAA0000A1Z5 &nbsp;|&nbsp; <b>Ph:</b> +91 9876543210", subtitle_style))
    elements.append(Spacer(1, 15))
    
    meta_data = [
        [Paragraph(f"<b>Invoice #:</b> {bill_id}", bold_style), Paragraph(f"<b>Payment:</b> {bill_data['payment_mode'].upper()}", bold_style)],
        [Paragraph(f"<b>Customer:</b> Walk-in", styles['Normal']), Paragraph(f"<b>Date:</b> Live System Record", styles['Normal'])]
    ]
    meta_table = Table(meta_data, colWidths=[260, 260])
    meta_table.setStyle(TableStyle([('BOTTOMPADDING', (0,0), (-1,-1), 4), ('LINEBELOW', (0,1), (-1,1), 1, colors.HexColor("#DDDDDD"))]))
    elements.append(meta_table)
    elements.append(Spacer(1, 15))
    
    table_data = [["Item Description", "Qty", "MRP (₹)", "Total (₹)"]]
    for item in items:
        table_data.append([str(item['sku_name']).title(), str(item['quantity']), f"{item['mrp']:.2f}", f"{item['total']:.2f}"])
        
    item_table = Table(table_data, colWidths=[240, 60, 100, 120])
    
    # Updated table to standard grey scale receipt style
    item_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#F2F2F2")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor("#111111")),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'), ('ALIGN', (0, 1), (0, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8), ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('LINEABOVE', (0, 0), (-1, 0), 2, colors.HexColor("#333333")),
        ('LINEBELOW', (0, 0), (-1, 0), 2, colors.HexColor("#333333")),
        ('LINEBELOW', (0, 1), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
    ]))
    elements.append(item_table)
    elements.append(Spacer(1, 10))
    
    totals_data = [
        ["Subtotal:", f"₹{bill_data['subtotal']:.2f}"],
        ["CGST:", f"₹{bill_data['cgst']:.2f}"],
        ["SGST:", f"₹{bill_data['sgst']:.2f}"],
        ["Grand Total:", f"₹{bill_data['total_amount']:.2f}"]
    ]
    totals_table = Table(totals_data, colWidths=[400, 120])
    totals_table.setStyle(TableStyle([
        ('ALIGN', (0, 0), (-1, -1), 'RIGHT'),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, -1), (-1, -1), 11),
        ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LINEABOVE', (0, -1), (-1, -1), 1.5, colors.HexColor("#111111")),
    ]))
    elements.append(totals_table)
    elements.append(Spacer(1, 30))
    
    footer_style = ParagraphStyle('Footer', parent=styles['Normal'], fontSize=8, textColor=colors.HexColor("#777777"), alignment=1)
    elements.append(Paragraph("Thank you for shopping with us! Goods once sold cannot be exchanged.", footer_style))
    doc.build(elements)
    return filename

def generate_pptx_deck(data):
    """Generates a multi-slide visual PowerPoint analysis deck."""
    os.makedirs("reports", exist_ok=True)
    filename = "reports/sales_analysis.pptx"
    prs = Presentation()
    
    # Slide 1: Executive Overview Chart
    slide1 = prs.slides.add_slide(prs.slide_layouts[5])
    title1 = slide1.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.5), Inches(0.8))
    title1.text_frame.text = "Kirana Store - Executive Sales Overview"
    
    chart_data = CategoryChartData()
    chart_data.categories = ['Orders', 'Revenue (₹)', 'GST (₹)']
    
    orders = float(data.get('total_orders') or 0)
    sales = float(data.get('total_sales') or 0.0)
    gst = float(data.get('total_gst') or 0.0)
    
    chart_data.add_series('Store Metrics', (orders, sales, gst))
    chart = slide1.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1.0), Inches(1.5), Inches(8.0), Inches(4.5), chart_data).chart
    chart.plots[0].has_data_labels = True
    
    # Slide 2: High/Low Velocity Comparison Slide (UPDATED ALIGNMENT)
    # Using layout 1 (Title and Content) instead of 5 for native alignment and bullets
    slide2 = prs.slides.add_slide(prs.slide_layouts[1]) 
    
    # Set Native Title
    title_shape = slide2.shapes.title
    title_shape.text = "Inventory Velocity: Best Sellers vs Stock Alerts"
    
    # Set Native Body Format
    body_shape = slide2.placeholders[1]
    tf = body_shape.text_frame
    tf.clear()  # Clear the default empty paragraph
    
    # Highest Selling Section (Level 0 - Header)
    p_high = tf.paragraphs[0]
    p_high.text = "🔥 Highest Selling Products (Demand Velocity):"
    p_high.font.bold = True
    p_high.font.size = Pt(22)
    p_high.level = 0
    
    # Highest Selling Data (Level 1 - Bullets)
    for ti in data.get('top_items', []):
        p_item = tf.add_paragraph()
        p_item.text = f"{ti['sku_name'].title()}: {ti['qty']} units sold"
        p_item.font.size = Pt(18)
        p_item.level = 1
        
    # Spacer Paragraph
    p_space = tf.add_paragraph()
    p_space.text = ""
    
    # Low Stock Section (Level 0 - Header)
    p_low = tf.add_paragraph()
    p_low.text = "⚠️ Low Stock / Reorder Watchlist:"
    p_low.font.bold = True
    p_low.font.size = Pt(22)
    p_low.level = 0
    
    # Low Stock Data (Level 1 - Bullets)
    for li in data.get('low_items', []):
        p_item_low = tf.add_paragraph()
        p_item_low.text = f"{li['sku_name'].title()}: {li['quantity']} remaining in stock"
        p_item_low.font.size = Pt(18)
        p_item_low.level = 1
        
    prs.save(filename)
    return filename